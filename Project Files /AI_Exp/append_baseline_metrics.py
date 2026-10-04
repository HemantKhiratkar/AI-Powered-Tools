from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from calculate_baseline_metrics import calculate_baseline_metrics


WORKBOOK_PATH = Path(__file__).resolve().parent / "T18_Supplier_scorecard_app.xlsx"
BASELINE_SHEET = "Baseline Metrics"
PERCENTAGE_COLUMNS = {"average_price_change", "overall_rejection_rate"}
HEADER_FILL = PatternFill(fill_type="solid", fgColor="173C36")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def append_baseline_metrics(workbook_path: Path) -> None:
    if not workbook_path.is_file():
        raise FileNotFoundError(
            f"Workbook not found: {workbook_path}. "
            "Place T18_Supplier_scorecard_app.xlsx next to this script."
        )

    suppliers = pd.read_excel(workbook_path, sheet_name="Suppliers")
    purchase_orders = pd.read_excel(workbook_path, sheet_name="Purchase Orders")
    metrics = calculate_baseline_metrics(suppliers, purchase_orders)

    # Replace only this report sheet on reruns; all other workbook sheets are retained.
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        metrics.to_excel(writer, sheet_name=BASELINE_SHEET, index=False)

    workbook = load_workbook(workbook_path)
    worksheet = workbook[BASELINE_SHEET]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        max_content_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = max_content_length + 2

    header_columns = {
        cell.value: cell.column for cell in worksheet[1] if cell.value is not None
    }
    for column_name in PERCENTAGE_COLUMNS & header_columns.keys():
        column_index = header_columns[column_name]
        for row in worksheet.iter_rows(
            min_row=2, min_col=column_index, max_col=column_index
        ):
            row[0].number_format = "0.00%"

    workbook.save(workbook_path)


if __name__ == "__main__":
    append_baseline_metrics(WORKBOOK_PATH)
    print(f"Appended styled '{BASELINE_SHEET}' sheet to {WORKBOOK_PATH}")
